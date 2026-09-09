"""Vista adapter: official investor workbook, with explicit units and provenance.

openpyxl is used only to READ source/user workbooks. The output author is
the separate Artifact Tool renderer.
"""
from __future__ import annotations

import hashlib
import math
import re
from datetime import datetime, timezone
from pathlib import Path
from urllib.request import Request, urlopen

import openpyxl

IR_URL = "https://www.vistaenergy.com/en/investors"
HISTORY_URL = "https://vista-energy.cdn.prismic.io/vista-energy/53GiiahTPNVSHrqM_20260828Vista-Historicalfinancialandoperationaldata.xlsx"
AUDIT_URLS = {
    2025: "https://vista-energy.cdn.prismic.io/vista-energy/afEJhMBOoF08xZTM_12.Vista-ConsolidatedFinancialStatements-VistaEnergySABdeCV-12.2025-withopinion-.pdf",
    2024: "https://vista-energy.cdn.prismic.io/vista-energy/aQq0Z7pReVYa4EQi_12.Vista-2024AuditedConsolidatedFinancialStatements.pdf",
}
ADR_URL = "https://vistaenergy.com/contenidos/VistaFY202420F.pdf"

# key, label, unit, source sheet, exact label prefix, occurrence, scale
# Expense outflows are standardized to positive values by negative scale.
METRICS = [
    ("revenue", "Reported revenue", "USD m", "Income Statement", "Revenue from contracts", 0, .001),
    ("oil_revenue", "Reported oil revenue", "USD m", "Income Statement", "Revenues from crude oil", 0, .001),
    ("gas_revenue", "Reported gas revenue", "USD m", "Income Statement", "Revenues from natural gas", 0, .001),
    ("ngl_revenue", "Reported NGL revenue", "USD m", "Income Statement", "Revenues from LPG", 0, .001),
    ("export_duties", "Export duties", "USD m", "Sales", "Export Duties", 0, 1),
    ("opex", "Lifting costs", "USD m", "Income Statement", "Operating costs", 0, -.001),
    ("royalties", "Royalties and other production taxes", "USD m", "Income Statement", "Royalties and others", 0, -.001),
    ("selling", "Reported selling expenses", "USD m", "Income Statement", "Selling expenses", 0, -.001),
    ("gna", "General and administrative expenses", "USD m", "Income Statement", "General and administrative", 0, -.001),
    ("dda", "Depreciation, depletion and amortization", "USD m", "Income Statement", "Depreciation, depletion", 0, -.001),
    ("ebit", "Reported operating profit", "USD m", "Income Statement", "Operating profit", 0, .001),
    ("net_income", "Net income", "USD m", "Income Statement", "Profit for the period, net", 0, .001),
    ("cash_taxes", "Cash income taxes paid", "USD m", "Cash Flow Statement", "Income tax payment", 0, -.001),
    ("cfo", "Cash flow from operations", "USD m", "Cash Flow Statement", "Net cash flows provided by operating", 0, .001),
    ("cash_capex", "Cash PPE and biological asset investment", "USD m", "Cash Flow Statement", "Payments for acquisitions of property", 0, -.001),
    ("intangible_capex", "Cash intangible investment", "USD m", "Cash Flow Statement", "Payments for acquisitions of other intangible", 0, -.001),
    ("lease_cash", "Lease cash payments", "USD m", "Cash Flow Statement", "Payment of lease", 0, -.001),
    ("acquisition_cash", "Cash business acquisitions", "USD m", "Cash Flow Statement", "Payments for Business Combination", 0, -.001),
    ("acquisition_gain", "Nonrecurring acquisition gain", "USD m", "Cash Flow Statement", "Gain from business combination", 0, -.001),
    ("cfi", "Cash flow from investing", "USD m", "Cash Flow Statement", "Net cash flows (used in) investing", 0, .001),
    ("cff", "Cash flow from financing", "USD m", "Cash Flow Statement", "Net cash flow provided by (used in) financing", 0, .001),
    ("opening_cash", "Opening cash and cash equivalents", "USD m", "Cash Flow Statement", "Cash and cash equivalents at beginning", 0, .001),
    ("fx_cash", "FX and other cash reconciliation", "USD m", "Cash Flow Statement", "Effect of exposure to changes", 0, .001),
    ("closing_cash", "Closing cash and cash equivalents", "USD m", "Cash Flow Statement", "Cash and cash equivalents at end", 0, .001),
    ("cash", "Cash and short-term investments", "USD m", "Balance Sheet", "Cash, bank balances", 0, .001),
    ("debt_nc", "Noncurrent borrowings", "USD m", "Balance Sheet", "Borrowings", 0, .001),
    ("debt_current", "Current borrowings", "USD m", "Balance Sheet", "Borrowings", 1, .001),
    ("lease_nc", "Noncurrent lease liabilities", "USD m", "Balance Sheet", "Lease liabilities", 0, .001),
    ("lease_current", "Current lease liabilities", "USD m", "Balance Sheet", "Lease liabilities", 1, .001),
    ("receivables", "Current trade and other receivables", "USD m", "Balance Sheet", "Trade and other receivables", 1, .001),
    ("inventory", "Inventories", "USD m", "Balance Sheet", "Inventories", 0, .001),
    ("payables", "Current trade and other payables", "USD m", "Balance Sheet", "Trade and other payables", 1, .001),
    ("deferred_payables", "Noncurrent trade and other payables", "USD m", "Balance Sheet", "Trade and other payables", 0, .001),
    ("associates", "Investments in associates", "USD m", "Balance Sheet", "Investments in associates", 0, .001),
    ("assets", "Total assets", "USD m", "Balance Sheet", "Total assets", 0, .001),
    ("liabilities", "Total liabilities", "USD m", "Balance Sheet", "Total liabilities", 0, .001),
    ("equity", "Total book equity", "USD m", "Balance Sheet", "Total equity", 0, .001),
    ("production", "Total production", "boe/day", "Production", "Total production by field", 0, 1),
    ("oil_production", "Oil production", "bbl/day", "Production", "Crude oil production by field", 0, 1),
    ("gas_production", "Gas production", "boe/day", "Production", "Natural Gas production by field", 0, 1),
    ("ngl_production", "NGL production", "boe/day", "Production", "NGL production by field", 0, 1),
    ("oil_sales", "Oil sales volume", "million bbl", "Sales", "Oil (MMbbl)", 0, 1),
    ("gas_sales", "Gas sales volume", "million MMBtu", "Sales", "Natural Gas (million", 0, 1),
    ("oil_price", "Realized oil price, net of duties", "USD/bbl", "Sales", "Oil ($/bbl)", 0, 1),
    ("gas_price", "Realized gas price", "USD/MMBtu", "Sales", "Natural Gas ($/MMBtu)", 0, 1),
    ("lifting_unit", "Reported lifting cost per boe", "USD/boe", "Opex & Capex", "Lifting Cost ($/boe)", 0, 1),
    ("ppe_additions", "Accrual PPE additions", "USD m", "Opex & Capex", "Capex ($MM)", 0, 1),
]


def download(url: str, target: Path) -> None:
    """Download to a temporary file first; never destroy a valid cache on failure."""
    request = Request(url, headers={"User-Agent": "VistaDCF/1.0 financial research"})
    with urlopen(request, timeout=60) as response:
        content = response.read()
    if not content:
        raise ValueError(f"Empty download: {url}")
    if target.suffix.lower()=='.xlsx' and not content.startswith(b'PK'):
        raise ValueError(f"Expected an XLSX file, received another content type: {url}")
    if target.suffix.lower()=='.pdf' and not content.startswith(b'%PDF'):
        raise ValueError(f"Expected a PDF file, received another content type: {url}")
    target.parent.mkdir(parents=True, exist_ok=True)
    temp = target.with_suffix(target.suffix + ".download")
    temp.write_bytes(content)
    temp.replace(target)


def discover_history_url() -> str:
    request = Request(IR_URL, headers={"User-Agent": "VistaDCF/1.0 financial research"})
    with urlopen(request, timeout=60) as response:
        html = response.read().decode("utf-8")
    urls = re.findall(r'https://[^\s"<>]+\.xlsx', html)
    urls = [u for u in urls if "historical" in u.lower()]
    if not urls:
        raise ValueError("Vista's historical workbook link was not found. Use --source-xlsx explicitly.")
    return urls[0].replace("&amp;", "&")


def annual_column(sheet, year: int) -> int:
    header_row = 7 if sheet.title in ("Balance Sheet", "Cash Flow Statement", "Production") else 6
    if sheet.title == "Balance Sheet":
        hits = [c.column for c in sheet[header_row] if re.fullmatch(
            rf"As of December 31, {year}", str(c.value).strip())]
    else:
        hits = [c.column for c in sheet[header_row] if str(c.value).strip() == str(year)]
    if len(hits) != 1:
        raise ValueError(f"Expected one annual column for {sheet.title} {year}, found {len(hits)}")
    return hits[0]


def read_number(value, location: str) -> float:
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value):
        raise ValueError(f"Missing/non-numeric source at {location}: {value!r}. Review the source or mapping.")
    return float(value)


def load_vista(path: Path, end_year: int = 2025, source_url: str = HISTORY_URL) -> dict:
    if end_year != 2025:
        raise ValueError("Vista v1 supports FY2023–2025. A new fiscal year requires reviewed audit and share metadata.")
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=False)
    history, facts = {}, []
    for year in range(end_year - 2, end_year + 1):
        values = {}
        for key, label, unit, sheet_name, prefix, occurrence, scale in METRICS:
            sheet = workbook[sheet_name]
            matches = [r for r in range(1, sheet.max_row + 1)
                       if str(sheet.cell(r, 2).value).strip().startswith(prefix)]
            if occurrence >= len(matches):
                raise ValueError(f"Missing source row: {sheet_name} / {prefix}")
            row, col = matches[occurrence], annual_column(sheet, year)
            cell = sheet.cell(row, col)
            value = read_number(cell.value, f"{sheet_name}!{cell.coordinate}") * scale
            values[key] = value
            facts.append(dict(key=key, label=label, year=year, value=value, unit=unit,
                              url=source_url, location=f"{sheet_name}!{cell.coordinate}",
                              raw_value=cell.value, scale=scale,
                              note="Company historical workbook, unaudited; audited statements prevail."))
        history[str(year)] = values
    workbook.close()
    return dict(company="Vista Energy", ticker="VIST", currency="USD", industry="upstream",
                domicile="Mexico", valuation_date="2025-12-31", history=history, facts=facts,
                source_url=source_url, source_sha256=hashlib.sha256(path.read_bytes()).hexdigest(),
                retrieved_at=datetime.fromtimestamp(path.stat().st_mtime,timezone.utc).isoformat(),
                parsed_at=datetime.now(timezone.utc).isoformat())


def audit_history(data: dict, raw_dir: Path) -> list[dict]:
    """Independent audited-PDF comparisons, selected by statement labels."""
    from pypdf import PdfReader
    patterns = {
        "revenue": r"Revenue from contracts with customers\s+5\s+([\d,]+)\s+([\d,]+)",
        "assets": r"Total assets\s+([\d,]+)\s+([\d,]+)",
        "liabilities": r"Total liabilities\s+([\d,]+)\s+([\d,]+)",
        "cfo": r"Net cash flows provided by operating activities\s+([\d,]+)\s+([\d,]+)",
        "closing_cash": r"Cash and cash equivalents at end of (?:year|period)\s+(?:20\s+|21\s+)?([\d,]+)\s+([\d,]+)",
    }
    checks = []
    for report_year in (2024, 2025):
        pdf = raw_dir / f"vista_{report_year}_audited.pdf"
        if not pdf.exists():
            raise ValueError(f"Audited report missing: {pdf}; run refresh or provide the cached report.")
        text = "\n".join(p.extract_text() for p in PdfReader(pdf).pages[:18])
        for key, pattern in patterns.items():
            match = re.search(pattern, text, re.IGNORECASE)
            if match is None:
                raise ValueError(f"Audit label not recognized in {report_year}: {key}; review changed PDF layout.")
            for year, group in ((report_year, 1), (report_year-1, 2)):
                if year == 2024 and report_year == 2024:
                    continue  # Prefer latest audited comparative.
                expected = float(match.group(group).replace(",", "")) / 1000
                actual = data["history"][str(year)][key]
                difference = actual - expected
                checks.append(dict(key=key, year=year, actual=actual, audited=expected,
                                   difference=difference, url=AUDIT_URLS[report_year]))
                if abs(difference) > .002:
                    raise ValueError(f"Audit mismatch {key} {year}: source {actual}, audited {expected}")
    return checks
